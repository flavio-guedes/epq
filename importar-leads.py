#!/usr/bin/env python3
"""
importar-leads.py
Importa leads a partir de um CSV para a planilha local
Inscricoes-Bolsao-EPQ-2026.xlsx.

Uso:
  python3 importar-leads.py caminho/do/arquivo.csv

CSV esperado (cabeçalho):
  nome,email,whatsapp,turma,data_envio

O script:
  - abre Inscricoes-Bolsao-EPQ-2026.xlsx na mesma pasta;
  - lê o CSV informado;
  - ignora e-mails já existentes na aba Inscricoes;
  - adiciona somente os leads novos;
  - preenche Origem = 'LP Bolsão' e Status do Contato = 'Novo';
  - salva a planilha.
"""

import csv
import os
import sys
from datetime import datetime
from openpyxl import load_workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
XLSX_PATH = os.path.join(BASE_DIR, "Inscricoes-Bolsao-EPQ-2026.xlsx")
SHEET_NAME = "Inscricoes"


def main(csv_path: str) -> int:
    if not os.path.isfile(csv_path):
        print(f"ERRO: arquivo não encontrado: {csv_path}")
        return 1

    if not os.path.isfile(XLSX_PATH):
        print(f"ERRO: planilha não encontrada: {XLSX_PATH}")
        return 1

    wb = load_workbook(XLSX_PATH)
    if SHEET_NAME not in wb.sheetnames:
        print(f"ERRO: aba '{SHEET_NAME}' não encontrada na planilha.")
        return 1

    ws = wb[SHEET_NAME]

    # Ler e-mails já existentes (coluna C = E-mail, índice 3)
    existing_emails = set()
    for row in ws.iter_rows(min_row=2, max_col=3, values_only=True):
        email = (row[2] or "").strip().lower()
        if email:
            existing_emails.add(email)

    # Ler CSV
    with open(csv_path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        rows = list(reader)

    if not rows:
        print("AVISO: CSV vazio ou sem dados.")
        return 0

    required = {"nome", "email", "whatsapp", "turma", "data_envio"}
    if not required.issubset({h.strip().lower() for h in reader.fieldnames or []}):
        missing = required - {h.strip().lower() for h in (reader.fieldnames or [])}
        print(f"ERRO: CSV está faltando colunas obrigatórias: {', '.join(sorted(missing))}")
        print(f"Esperado: {', '.join(sorted(required))}")
        return 1

    new_rows = []
    duplicates = 0
    errors = 0

    for row in rows:
        nome = (row.get("nome") or "").strip()
        email = (row.get("email") or "").strip().lower()
        whatsapp = (row.get("whatsapp") or "").strip()
        turma = (row.get("turma") or "").strip()
        data_envio = (row.get("data_envio") or "").strip()

        if not email:
            errors += 1
            continue

        if email in existing_emails:
            duplicates += 1
            continue

        new_rows.append((nome, email, whatsapp, turma, data_envio))
        existing_emails.add(email)

    # Determinar próxima linha disponível
    next_row = ws.max_row + 1
    if next_row == 1:
        next_row = 2

    for nome, email, whatsapp, turma, data_envio in new_rows:
        ws.cell(row=next_row, column=1, value=datetime.now())
        ws.cell(row=next_row, column=2, value=nome)
        ws.cell(row=next_row, column=3, value=email)
        ws.cell(row=next_row, column=4, value=whatsapp)
        ws.cell(row=next_row, column=5, value=turma)
        ws.cell(row=next_row, column=6, value="LP Bolsão")
        ws.cell(row=next_row, column=7, value="Novo")
        # colunas 8 a 11 ficam vazias
        next_row += 1

    wb.save(XLSX_PATH)

    print(f"Arquivo CSV: {csv_path}")
    print(f"Registros lidos: {len(rows)}")
    print(f"Novos registros importados: {len(new_rows)}")
    print(f"Duplicados ignorados: {duplicates}")
    print(f"Erros: {errors}")
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Uso: python3 importar-leads.py caminho/do/arquivo.csv")
        sys.exit(1)
    sys.exit(main(sys.argv[1]))
